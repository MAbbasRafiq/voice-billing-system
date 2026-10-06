# -*- coding: utf-8 -*-
"""Fill missing items.urdu_name (DB + Excel) with Urdu-script transliteration of English names."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.database.connection import get_connection  # noqa: E402
from app.database.queries import invalidate_items_cache  # noqa: E402

DB_EXCEL = ROOT / "data" / "C_P_LIST.xlsx"

# Phrase replacements (longest first) — English UPPER → Urdu
PHRASES = [
    ("AIR FILTER", "ایئر فلٹر"),
    ("CHAIN KIT", "چین کٹ"),
    ("WHEEL CHAIN", "ویل چین"),
    ("HANDLE T ONLY", "ہینڈل ٹی اونلی"),
    ("HANDLE CONE DUST COVER", "ہینڈل کون ڈسٹ کور"),
    ("HANDLE CONE SET", "ہینڈل کون سیٹ"),
    ("HANDLE BAR", "ہینڈل بار"),
    ("HANDLE CONE", "ہینڈل کون"),
    ("BACK LIGHT", "بیک لائٹ"),
    ("BACK TACK", "بیک ٹیک"),
    ("TIMING CHAIN", "ٹائمنگ چین"),
    ("CHAIN LOCK", "چین لاک"),
    ("CHAIN COVER", "چین کور"),
    ("ALLOY RIM FRONT", "ایلوی رم فرنٹ"),
    ("ALLOY RIM REAR", "ایلوی رم ریئر"),
    ("C.D.I UNIT", "سی ڈی آئی یونٹ"),
    ("CDI UNIT", "سی ڈی آئی یونٹ"),
    ("BATTERY BOX", "بیٹری باکس"),
    ("NUMBER PLATE", "نمبر پلیٹ"),
    ("SIDE MIRROR", "سائیڈ مرر"),
    ("HEAD LIGHT", "ہیڈ لائٹ"),
    ("TAIL LIGHT", "ٹیل لائٹ"),
    ("STOP LIGHT", "اسٹاپ لائٹ"),
    ("INDICATOR LIGHT", "انڈیکیٹر لائٹ"),
    ("BRAKE SHOE", "بریک شو"),
    ("BRAKE PAD", "بریک پیڈ"),
    ("DISC PAD", "ڈسک پیڈ"),
    ("CLUTCH PLATE", "کلچ پلیٹ"),
    ("OIL FILTER", "آئل فلٹر"),
    ("OIL SEAL", "آئل سیل"),
    ("O RING", "او رنگ"),
    ("SPARK PLUG", "سپارک پلگ"),
    ("FUEL COCK", "فیول کاک"),
    ("FUEL PIPE", "فیول پائپ"),
    ("ACCELERATOR CABLE", "ایکسلریٹر کیبل"),
    ("CLUTCH CABLE", "کلچ کیبل"),
    ("BRAKE CABLE", "بریک کیبل"),
    ("SPEEDO CABLE", "سپیڈو کیبل"),
    ("THROTTLE CABLE", "تھروٹل کیبل"),
    ("KICK STARTER", "کک اسٹارٹر"),
    ("SELF STARTER", "سیلف اسٹارٹر"),
    ("SHOCK ABSORBER", "شاک ابزاربر"),
    ("REAR VIEW", "ریئر ویو"),
    ("FRONT FORK", "فرنٹ فورک"),
    ("STEERING LOCK", "اسٹیئرنگ لاک"),
    ("IGNITION COIL", "اگنشن کوائل"),
    ("CONDENSER", "کنڈینسر"),
    ("RECTIFIER", "ریکٹیفائر"),
    ("REGULATOR", "ریگولیٹر"),
    ("CARBURETOR", "کاربریٹر"),
    ("CARBURETTOR", "کاربریٹر"),
    ("AIR CLEANER", "ایئر کلینر"),
    ("SILENCER", "سائیلنسر"),
    ("EXHAUST PIPE", "ایگزاسٹ پائپ"),
    ("MUFFLER", "مفلر"),
    ("PACKING", "پیکنگ"),
    ("GASKET", "گیسکٹ"),
]

WORD = {
    "AIR": "ایئر",
    "FILTER": "فلٹر",
    "FOAM": "فوم",
    "BODY": "باڈی",
    "RUBBER": "ربڑ",
    "BOX": "باکس",
    "CORE": "کور",
    "COMPLETE": "کمپلیٹ",
    "COMP": "کمپ",
    "CHAIN": "چین",
    "KIT": "کٹ",
    "WHEEL": "ویل",
    "HANDLE": "ہینڈل",
    "BAR": "بار",
    "CONE": "کون",
    "DUST": "ڈسٹ",
    "COVER": "کور",
    "SET": "سیٹ",
    "BACK": "بیک",
    "LIGHT": "لائٹ",
    "FRONT": "فرنٹ",
    "REAR": "ریئر",
    "LOCK": "لاک",
    "TIMING": "ٹائمنگ",
    "ALLOY": "ایلوی",
    "RIM": "رم",
    "ONLY": "اونلی",
    "BRIDGE": "برج",
    "ROD": "راڈ",
    "PLATE": "پلیٹ",
    "UNIT": "یونٹ",
    "LEED": "لیڈ",
    "LEAD": "لیڈ",
    "LENS": "لینز",
    "LED": "ایل ای ڈی",
    "SWITCH": "سوئچ",
    "CABLE": "کیبل",
    "BRAKE": "بریک",
    "CLUTCH": "کلچ",
    "GEAR": "گیئر",
    "SHOCK": "شاک",
    "ABSORBER": "ابزاربر",
    "SPRING": "سپرنگ",
    "BOLT": "بولٹ",
    "NUT": "نٹ",
    "WASHER": "واشر",
    "BEARING": "بیرنگ",
    "BUSH": "بش",
    "PIN": "پن",
    "PINS": "پنز",
    "BATTERY": "بیٹری",
    "HORN": "ہارن",
    "MIRROR": "مرر",
    "SPEEDO": "سپیڈو",
    "METER": "میٹر",
    "TANK": "ٹینک",
    "SEAT": "سیٹ",
    "STAND": "اسٹینڈ",
    "SIDE": "سائیڈ",
    "MAIN": "مین",
    "OIL": "آئل",
    "PUMP": "پمپ",
    "RING": "رنگ",
    "PISTON": "پسٹن",
    "VALVE": "والو",
    "CAM": "کیم",
    "SHAFT": "شافٹ",
    "SPROCKET": "اسپراکیٹ",
    "DISC": "ڈسک",
    "PAD": "پیڈ",
    "PADS": "پیڈز",
    "SHOES": "شوز",
    "SHOE": "شو",
    "DRUM": "ڈرم",
    "HUB": "ہب",
    "SPOKE": "سپوک",
    "TUBE": "ٹیوب",
    "TYRE": "ٹائر",
    "TIRE": "ٹائر",
    "GRIP": "گرپ",
    "LEVER": "لیور",
    "PEDAL": "پیڈل",
    "KICK": "کک",
    "STARTER": "اسٹارٹر",
    "RELAY": "ریلے",
    "FUSE": "فیوز",
    "WIRE": "وائر",
    "HARNESS": "ہارنس",
    "SOCKET": "ساکٹ",
    "BULB": "بلب",
    "FLASHER": "فلشر",
    "INDICATOR": "انڈیکیٹر",
    "TAIL": "ٹیل",
    "HEAD": "ہیڈ",
    "STOP": "اسٹاپ",
    "NUMBER": "نمبر",
    "GUARD": "گارڈ",
    "FENDER": "فینڈر",
    "MUDGUARD": "مڈگارڈ",
    "SILENCER": "سائیلنسر",
    "EXHAUST": "ایگزاسٹ",
    "GASKET": "گیسکٹ",
    "PACKING": "پیکنگ",
    "SEAL": "سیل",
    "WITH": "ودھ",
    "WITHOUT": "ودھاؤٹ",
    "AND": "اینڈ",
    "FOR": "فار",
    "TYPE": "ٹائپ",
    "SIZE": "سائز",
    "BLACK": "بلیک",
    "WHITE": "وائٹ",
    "RED": "ریڈ",
    "BLUE": "بلیو",
    "GREEN": "گرین",
    "YELLOW": "یلو",
    "CHROME": "کروم",
    "STEEL": "سٹیل",
    "ALUMINIUM": "ایلومینیم",
    "ALUMINUM": "ایلومینیم",
    "PLASTIC": "پلاسٹک",
    "ORIGINAL": "اوریجنل",
    "LOCAL": "لوکل",
    "BOTTLE": "بوتل",
    "HOSE": "ہوز",
    "PIPE": "پائپ",
    "WHEELER": "وہیلر",
    "CAP": "کیپ",
    "CLIP": "کلپ",
    "CLAMP": "کلیمپ",
    "JOINT": "جوائنٹ",
    "ARM": "آرم",
    "BRACKET": "بریکٹ",
    "HOLDER": "ہولڈر",
    "MOUNT": "ماؤنٹ",
    "MOUNTING": "ماؤنٹنگ",
    "ASSY": "ایسمبلی",
    "ASSEMBLY": "ایسمبلی",
    "KIT": "کٹ",
    "SET": "سیٹ",
    "RH": "دائیں",
    "LH": "بائیں",
    "RIGHT": "رائٹ",
    "LEFT": "لیفٹ",
    "UPPER": "اپر",
    "LOWER": "لوئر",
    "INNER": "انر",
    "OUTER": "آؤٹر",
    "TOP": "ٹاپ",
    "BOTTOM": "باٹم",
    "NEW": "نیو",
    "OLD": "اولڈ",
    "MODEL": "ماڈل",
    "COLOR": "کلر",
    "COLOUR": "کلر",
    "ENGINE": "انجین",
    "MOTOR": "موٹر",
    "ELECTRIC": "الیکٹرک",
    "EV": "ای وی",
    "SWITCH": "سوئچ",
    "KEY": "کی",
    "LOCK": "لاک",
    "SET": "سیٹ",
    "PAIR": "پیئر",
    "PCS": "پیسز",
    "PC": "پیس",
    "BASE": "بیس",
    "CENTER": "سینٹر",
    "CENTRE": "سینٹر",
    "AXLE": "ایکسل",
    "ARMATURE": "آرمیچر",
    "FUEL": "فیول",
    "TANK": "ٹینک",
    "REFLECTOR": "ریفلکٹر",
    "PLASTIC": "پلاسٹک",
    "COMPLETE": "کمپلیٹ",
    "ASSY": "ایسمبلی",
    "ASSEMBLY": "ایسمبلی",
    "MT": "MT",
    "RH": "RH",
    "LH": "LH",
    "FULL": "فل",
    "SMALL": "سمال",
    "LARGE": "لارج",
    "BIG": "بگ",
    "LONG": "لانگ",
    "SHORT": "شارٹ",
    "HEAVY": "ہیوی",
    "LIGHT": "لائٹ",
    "DUTY": "ڈیوٹی",
    "STANDARD": "سٹینڈرڈ",
    "SPECIAL": "سپیشل",
    "UNIVERSAL": "یونیورسل",
    "ADJUSTABLE": "ایڈجسٹیبل",
    "DOUBLE": "ڈبل",
    "SINGLE": "سنگل",
    "TRIPLE": "ٹرپل",
    "FOUR": "فور",
    "FIVE": "فائیو",
    "SIX": "سکس",
    "PLUS": "پلس",
    "PRO": "پرو",
    "MAX": "میکس",
    "MINI": "منی",
    "SUPER": "سپر",
    "ULTRA": "الٹرا",
    "PREMIUM": "پریمیم",
    "ECONOMY": "اکانومی",
    "GENUINE": "جینیوئن",
    "COPY": "کاپی",
    "HIGH": "ہائی",
    "LOW": "لو",
    "PRESSURE": "پریشر",
    "TEMPERATURE": "ٹمپریچر",
    "SENSOR": "سینسر",
    "SWITCH": "سوئچ",
    "BUTTON": "بٹن",
    "KNOB": "ناب",
    "DIAL": "ڈائل",
    "GAUGE": "گیج",
    "CLUSTER": "کلسٹر",
    "PANEL": "پینل",
    "DASHBOARD": "ڈیش بورڈ",
    "FAIRING": "فیئرنگ",
    "COWL": "کاؤل",
    "VISOR": "وائزر",
    "SCREEN": "اسکرین",
    "GLASS": "گلاس",
    "REFLECTOR": "ریفلکٹر",
    "EMBLEM": "ایمبلم",
    "STICKER": "اسٹکر",
    "DECAL": "ڈیکل",
    "TRIM": "ٹرم",
    "MOULDING": "مولڈنگ",
    "MOLDING": "مولڈنگ",
    "BEADING": "بیڈنگ",
    "STRIP": "سٹرپ",
    "TAPE": "ٹیپ",
    "FOAM": "فوم",
    "SPONGE": "سپنج",
    "FELT": "فیلٹ",
    "CORK": "کارک",
    "PAPER": "پیپر",
    "METAL": "میٹل",
    "NYLON": "نایلان",
    "TEFLON": "ٹیفلان",
    "COPPER": "کاپر",
    "BRASS": "براس",
    "IRON": "آئرن",
    "CAST": "کاسٹ",
    "FORGED": "فورجد",
    "MACHINED": "مشینڈ",
    "WELDED": "ویلڈڈ",
    "CHROME": "کروم",
    "NICKEL": "نکل",
    "ZINC": "زنک",
    "POWDER": "پاؤڈر",
    "COATED": "کوٹڈ",
    "PAINTED": "پینٹڈ",
    "POLISHED": "پالشڈ",
    "MATTE": "میٹ",
    "GLOSSY": "گلاسی",
    "TRANSPARENT": "ٹرانسپرنٹ",
    "CLEAR": "کلئیر",
    "SMOKE": "سموک",
    "TINTED": "ٹنٹڈ",
    "AMBER": "ایمبر",
    "ORANGE": "اورنج",
    "PURPLE": "پرپل",
    "PINK": "پنک",
    "GREY": "گرے",
    "GRAY": "گرے",
    "SILVER": "سلور",
    "GOLD": "گولڈ",
    "BROWN": "براؤن",
    "W/O": "ودھاؤٹ",
    "WO": "ودھاؤٹ",
    "N": "این",
    "T": "ٹی",
    "H": "ایچ",
    "L": "ایل",
    "R": "آر",
    "S": "ایس",
    "A": "اے",
    "B": "بی",
    "C": "سی",
    "D": "ڈی",
    "E": "ای",
    "F": "ایف",
    "G": "جی",
    "I": "آئی",
    "J": "جے",
    "K": "کے",
    "M": "ایم",
    "O": "او",
    "P": "پی",
    "Q": "کیو",
    "U": "یو",
    "V": "وی",
    "W": "ڈبلیو",
    "X": "ایکس",
    "Y": "وائے",
    "Z": "زیڈ",
}

# Simple roman → Urdu for unknown whole words (shop-style phonetic)
_DIGRAPHS = [
    ("SCH", "ش"),
    ("TCH", "چ"),
    ("CH", "چ"),
    ("SH", "ش"),
    ("TH", "تھ"),
    ("KH", "کھ"),
    ("GH", "گھ"),
    ("PH", "ف"),
    ("QU", "کو"),
    ("CK", "ک"),
    ("EE", "ی"),
    ("OO", "و"),
    ("AI", "ای"),
    ("AY", "ے"),
    ("OA", "و"),
    ("OU", "او"),
    ("OW", "او"),
    ("EA", "ی"),
    ("IE", "ی"),
    ("UE", "یو"),
    ("NG", "نگ"),
]


_LETTERS = {
    "A": "ا",
    "B": "ب",
    "C": "ک",
    "D": "ڈ",
    "E": "ی",
    "F": "ف",
    "G": "گ",
    "H": "ہ",
    "I": "ی",
    "J": "ج",
    "K": "ک",
    "L": "ل",
    "M": "م",
    "N": "ن",
    "O": "و",
    "P": "پ",
    "Q": "ق",
    "R": "ر",
    "S": "س",
    "T": "ٹ",
    "U": "و",
    "V": "و",
    "W": "و",
    "X": "کس",
    "Y": "ی",
    "Z": "ز",
}


def _phonetic_word(word: str) -> str:
    w = word.upper()
    if not w.isalpha():
        return word
    if w in WORD:
        return WORD[w]
    out: list[str] = []
    i = 0
    while i < len(w):
        matched = False
        for eng, ur in _DIGRAPHS:
            if w.startswith(eng, i):
                out.append(ur)
                i += len(eng)
                matched = True
                break
        if matched:
            continue
        out.append(_LETTERS.get(w[i], w[i]))
        i += 1
    return "".join(out)


def _is_code_token(tok: str) -> bool:
    """Keep item codes / sizes / model-like tokens in Latin."""
    if re.search(r"\d", tok):
        return True
    if re.fullmatch(r"[A-Z]{1,3}\d+[A-Z0-9.\-]*", tok, re.I):
        return True
    if re.fullmatch(r"\d+[A-Z0-9.\-/X]*", tok, re.I):
        return True
    # Short all-caps codes (MT, LH already mapped when known)
    if re.fullmatch(r"[A-Z]{2,4}", tok) and tok.upper() not in WORD:
        return True
    return False


def transliterate_name(name: str) -> str:
    raw = (name or "").strip()
    if not raw:
        return ""

    work = raw.upper()
    # Protect phrase hits with placeholders
    protected: list[str] = []
    for eng, ur in sorted(PHRASES, key=lambda x: -len(x[0])):
        if eng in work:
            placeholder = f"«P{len(protected)}»"
            work = work.replace(eng, placeholder, 1)
            protected.append(ur)

    tokens = re.findall(r"[A-Z]+(?:\.[A-Z]+)*|\d+[A-Z0-9.\-/X]*|[^\sA-Z0-9]+|\s+", work, flags=re.I)
    # Simpler split
    pieces = re.split(r"(\s+|[()\[\]{},/]+)", work)
    out: list[str] = []
    for p in pieces:
        if not p:
            continue
        m = re.fullmatch(r"«P(\d+)»", p)
        if m:
            out.append(protected[int(m.group(1))])
            continue
        if re.fullmatch(r"\s+|[()\[\]{},/]+", p):
            out.append(p)
            continue
        # restore case-less token
        clean = p.replace(".", "")
        if p.startswith("«"):
            out.append(p)
            continue
        if _is_code_token(p):
            out.append(p)
            continue
        key = p.upper().replace(".", "")
        if key in WORD:
            out.append(WORD[key])
        elif len(key) == 1 and key in WORD:
            out.append(WORD[key])
        else:
            out.append(_phonetic_word(key) if key.isalpha() else p)

    text = "".join(out)
    text = re.sub(r"\s+", " ", text).strip()
    # Fix any leftover placeholders
    for i, ur in enumerate(protected):
        text = text.replace(f"«P{i}»", ur)
    return text


def main() -> None:
    conn = get_connection()
    missing = conn.execute(
        """
        SELECT id, name, urdu_name FROM items
        WHERE urdu_name IS NULL OR trim(urdu_name) IN ('', '-', '—', '.', '–')
        ORDER BY name, id
        """
    ).fetchall()
    print(f"missing rows: {len(missing)}")

    existing = conn.execute(
        """
        SELECT name, urdu_name FROM items
        WHERE urdu_name IS NOT NULL AND trim(urdu_name) NOT IN ('', '-', '—', '.', '–')
        """
    ).fetchall()
    by_name: dict[str, str] = {}
    for r in existing:
        key = (r["name"] or "").strip().upper()
        if key and key not in by_name:
            by_name[key] = (r["urdu_name"] or "").strip()

    cache = dict(by_name)
    updates: list[tuple[str, int]] = []
    for r in missing:
        key = (r["name"] or "").strip().upper()
        if not key:
            continue
        if key not in cache:
            cache[key] = transliterate_name(r["name"] or "")
        ur = cache[key]
        if ur:
            updates.append((ur, int(r["id"])))

    print(f"updating: {len(updates)}")
    # UTF-8 samples to file to avoid Windows console issues
    sample_path = ROOT / "scripts" / "_urdu_fill_sample.txt"
    with sample_path.open("w", encoding="utf-8") as f:
        for ur, iid in updates[:40]:
            name = conn.execute("SELECT name FROM items WHERE id=?", (iid,)).fetchone()["name"]
            f.write(f"{name} => {ur}\n")
    print(f"wrote sample: {sample_path}")

    conn.executemany("UPDATE items SET urdu_name = ? WHERE id = ?", updates)
    conn.commit()
    left = conn.execute(
        """
        SELECT count(1) AS n FROM items
        WHERE urdu_name IS NULL OR trim(urdu_name) IN ('', '-', '—', '.', '–')
        """
    ).fetchone()["n"]
    print(f"remaining missing in DB: {left}")
    conn.close()
    invalidate_items_cache()

    # Excel update
    if not DB_EXCEL.exists():
        print("Excel missing; DB updated only")
        return

    wb = load_workbook(DB_EXCEL)
    cells = 0
    for ws in wb.worksheets:
        header_row = name_col = urdu_col = None
        for r in range(1, min(20, ws.max_row or 1) + 1):
            vals = [ws.cell(r, c).value for c in range(1, min(25, (ws.max_column or 1) + 1))]
            norms = [
                " ".join(str(v).strip().lower().split()) if v is not None else "" for v in vals
            ]
            if "name" not in norms:
                continue
            header_row = r
            name_col = norms.index("name") + 1
            if "urdu name" in norms:
                urdu_col = norms.index("urdu name") + 1
            else:
                urdu_col = name_col + 1
            break
        if not header_row or not name_col or not urdu_col:
            continue
        for r in range(header_row + 1, (ws.max_row or 0) + 1):
            name = ws.cell(r, name_col).value
            if not name or not str(name).strip():
                continue
            key = str(name).strip().upper()
            ur = cache.get(key)
            if not ur:
                ur = transliterate_name(str(name))
                cache[key] = ur
            if not ur:
                continue
            cell = ws.cell(r, urdu_col)
            cur = "" if cell.value is None else str(cell.value).strip()
            if cur and cur not in {"-", "—", "–", "."}:
                continue
            cell.value = ur
            cells += 1
    wb.save(DB_EXCEL)
    print(f"excel cells filled: {cells}")
    print("done")


if __name__ == "__main__":
    main()
